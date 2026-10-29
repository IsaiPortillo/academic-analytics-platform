"""Lectura de archivos subidos (Excel o CSV) para las importaciones masivas.

Cada importación sigue el mismo recorrido: el usuario descarga una plantilla
generada por el sistema, la llena, la sube y ve una VISTA PREVIA con lo que va
a cambiar (filas nuevas, valores que se sobrescriben y errores). Nada se
escribe hasta que confirma. Si hay un solo error, no se puede confirmar: un
archivo a medias deja la base en un estado que nadie revisó.

Este módulo solo convierte el archivo en filas de texto; qué significa cada
columna lo decide cada módulo (notas, asistencia, matrícula).
"""

import csv
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from fastapi import UploadFile

TAMANO_MAXIMO = 2 * 1024 * 1024   # 2 MB: una plantilla llena pesa unos pocos KB
FILAS_MAXIMAS = 5000
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ErrorArchivo(ValueError):
    """El archivo no se puede leer (formato, tamaño, sin cabecera)."""


def normalizar(texto: Any) -> str:
    """Para comparar cabeceras: sin acentos, minúsculas, espacios simples."""
    limpio = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", limpio).strip().lower()


async def leer_filas(archivo: UploadFile) -> list[list[Any]]:
    contenido = await archivo.read(TAMANO_MAXIMO + 1)
    if not contenido:
        raise ErrorArchivo("El archivo está vacío.")
    if len(contenido) > TAMANO_MAXIMO:
        raise ErrorArchivo("El archivo supera 2 MB.")
    nombre = (archivo.filename or "").lower()
    if nombre.endswith(".xlsx") or contenido[:2] == b"PK":
        filas = _leer_xlsx(contenido)
    elif nombre.endswith(".csv") or nombre.endswith(".txt"):
        filas = _leer_csv(contenido)
    else:
        raise ErrorArchivo("Formato no reconocido: sube la plantilla en Excel (.xlsx) o CSV.")
    filas = [f for f in filas if any(c not in (None, "") for c in f)]
    if len(filas) > FILAS_MAXIMAS:
        raise ErrorArchivo(f"El archivo tiene más de {FILAS_MAXIMAS:,} filas.")
    return filas


def _leer_xlsx(contenido: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook  # noqa: PLC0415
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - cualquier fallo de lectura es "no es un xlsx válido"
        raise ErrorArchivo("No se pudo leer el Excel. ¿Es la plantilla descargada del sistema?") from exc
    hoja = libro.worksheets[0]
    filas = [list(f) for f in hoja.iter_rows(values_only=True, max_row=FILAS_MAXIMAS + 20)]
    libro.close()
    return filas


def _leer_csv(contenido: bytes) -> list[list[Any]]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")
    # Excel en español guarda los CSV con punto y coma.
    delimitador = ";" if texto.split("\n", 1)[0].count(";") > texto.split("\n", 1)[0].count(",") else ","
    return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def ubicar_cabecera(filas: list[list[Any]], primera_columna: str) -> int:
    """Índice de la fila de cabecera: la primera cuya celda A es `primera_columna`.
    La plantilla lleva arriba un título y una línea de instrucciones."""
    buscada = normalizar(primera_columna)
    for i, fila in enumerate(filas[:15]):
        if fila and normalizar(fila[0]) == buscada:
            return i
    raise ErrorArchivo(f"No se encontró la cabecera «{primera_columna}». Usa la plantilla descargada del sistema.")


def como_entero(valor: Any) -> Optional[int]:
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int,)):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else None
    texto = str(valor).strip()
    return int(texto) if re.fullmatch(r"\d+", texto) else None


def como_nota(valor: Any) -> Optional[Decimal]:
    """Nota de 0 a 10 con dos decimales. Acepta coma decimal. None si no es válida."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        numero = Decimal(str(valor))
    else:
        try:
            numero = Decimal(str(valor).strip().replace(",", "."))
        except InvalidOperation:
            return None
    if not numero.is_finite() or numero < 0 or numero > 10:
        return None
    return numero.quantize(Decimal("0.01"))


def como_fecha(valor: Any) -> Optional[date]:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def vacio(valor: Any) -> bool:
    return valor is None or str(valor).strip() == ""


def plantilla_xlsx(titulo: str, instrucciones: str, cabeceras: list[str], filas: list[list[Any]],
                   validaciones: Optional[dict[int, tuple]] = None, anchos: Optional[list[int]] = None) -> bytes:
    """Plantilla para llenar y volver a subir.

    Fila 1 título, fila 2 instrucciones, fila 4 cabecera, datos desde la 5. La
    columna A (el identificador) va en gris: el sistema la usa para ubicar cada
    fila y no se debe cambiar. `validaciones` restringe lo que Excel deja
    escribir en una columna: {índice: ("decimal", mínimo, máximo)} o
    {índice: ("lista", [valores])}.
    """
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Border, Font, PatternFill, Side  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415
    from openpyxl.worksheet.datavalidation import DataValidation  # noqa: PLC0415

    from .exportar import COLOR_MARCA, COLOR_TENUE  # noqa: PLC0415

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja["A1"] = titulo
    hoja["A1"].font = Font(bold=True, size=13, color=COLOR_MARCA)
    hoja["A2"] = instrucciones
    hoja["A2"].font = Font(size=9, color=COLOR_TENUE)
    fila_cabecera = 4
    borde = Border(bottom=Side(style="medium", color=COLOR_MARCA))
    gris = PatternFill("solid", fgColor="EEF1F5")
    for j, cabecera in enumerate(cabeceras, start=1):
        celda = hoja.cell(row=fila_cabecera, column=j, value=cabecera)
        celda.font = Font(bold=True)
        celda.border = borde
        hoja.column_dimensions[get_column_letter(j)].width = (anchos[j - 1] if anchos else
                                                              max(len(str(cabecera)) + 4, 14))
    for i, fila in enumerate(filas, start=fila_cabecera + 1):
        for j, valor in enumerate(fila, start=1):
            celda = hoja.cell(row=i, column=j, value=valor)
            if j == 1:
                celda.fill = gris
    hoja.freeze_panes = f"B{fila_cabecera + 1}"

    ultima = fila_cabecera + max(len(filas), 1)
    for indice, regla in (validaciones or {}).items():
        letra = get_column_letter(indice + 1)
        if regla[0] == "decimal":
            dv = DataValidation(type="decimal", operator="between", formula1=str(regla[1]),
                                formula2=str(regla[2]), allow_blank=True)
            dv.error = f"Debe ser un número entre {regla[1]} y {regla[2]}."
        else:
            dv = DataValidation(type="list", formula1='"' + ",".join(regla[1]) + '"', allow_blank=True)
            dv.error = "Elige uno de los valores de la lista."
        dv.showErrorMessage = True
        hoja.add_data_validation(dv)
        dv.add(f"{letra}{fila_cabecera + 1}:{letra}{ultima + 200}")

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
