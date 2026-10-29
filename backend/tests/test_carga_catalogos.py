"""Carga masiva de catálogos: lectura de la plantilla, vista previa y confirmación.

    cd backend && python -m unittest discover -s tests -v

La lógica (analizar, empaquetar, leer_carga) no toca la base; aplicar_carga se prueba
con una sesión de doble.
"""

import io
import json
import os
import unittest
from dataclasses import replace
from datetime import date, datetime
from types import SimpleNamespace
from unittest import mock

# Las rutas importan la configuración, que exige estas dos variables. Valores de
# prueba: aquí no se conecta a nada.
os.environ.setdefault("APP_DB_PASSWORD", "x")
os.environ.setdefault("APP_SECRET_KEY", "y")

from app import carga_catalogos as carga  # noqa: E402
from app.catalogos import CATALOGOS, hash_carnet  # noqa: E402
from app.importar import ErrorArchivo  # noqa: E402

MATERIAS = CATALOGOS["materias"]
PERIODOS = CATALOGOS["periodos"]
SECCIONES = CATALOGOS["secciones"]
ESTUDIANTES = CATALOGOS["estudiantes"]
REFS = {"carrera_id": {"SIS": 1, "ADM": 2}, "periodo_id": {"2026-I": 7}, "materia_id": {"MAT101": 11},
        "docente_id": {"DOC-1": 5}}


def _con_cabecera(catalogo, *filas):
    return [["Carga masiva"], ["instrucciones"], [], carga.cabeceras(catalogo), *filas]


def _materia(codigo="MAT101", nombre="Matemática I", uv=4, ciclo=1, carrera="SIS", activo="Sí"):
    return [codigo, nombre, uv, ciclo, carrera, activo]


class TestLecturaDeMaterias(unittest.TestCase):
    def _analizar(self, filas, existentes=None):
        return carga.analizar(MATERIAS, _con_cabecera(MATERIAS, *filas), REFS, existentes or {})

    def test_una_fila_valida_es_una_materia_nueva(self):
        cambios, errores = self._analizar([_materia()])
        self.assertEqual(errores, [])
        self.assertEqual([c.tipo for c in cambios], ["nueva"])
        valores = cambios[0].valores
        self.assertEqual((valores["codigo_materia"], valores["carrera_id"], valores["unidades_valorativas"], valores["activo"]),
                         ("MAT101", 1, 4, True))

    def test_el_codigo_se_acepta_con_o_sin_nombre_y_en_minusculas(self):
        cambios, errores = self._analizar([_materia(carrera="sis · Ingeniería en Sistemas")])
        self.assertEqual((errores, cambios[0].valores["carrera_id"]), ([], 1))

    def test_activo_vacio_cuenta_como_si_y_trabaja_vacio_como_no(self):
        cambios, _ = self._analizar([_materia(activo=None)])
        self.assertTrue(cambios[0].valores["activo"])
        cambios, _ = carga.analizar(ESTUDIANTES, _con_cabecera(ESTUDIANTES, ["A1", 2024, "SIS", None, "REGULAR", None]),
                                    REFS, {})
        self.assertFalse(cambios[0].valores["trabaja"])
        self.assertTrue(cambios[0].valores["activo"])

    def test_un_numero_de_excel_con_decimales_cero_es_entero(self):
        cambios, errores = self._analizar([_materia(uv=4.0, ciclo=2.0)])
        self.assertEqual((errores, cambios[0].valores["unidades_valorativas"], cambios[0].valores["ciclo_plan"]), ([], 4, 2))

    def test_distingue_nueva_cambio_e_igual_contra_lo_existente(self):
        existentes = {
            ("MAT101",): {"_id": 31, "codigo_materia": "MAT101", "nombre": "Matemática I", "unidades_valorativas": 4,
                          "ciclo_plan": 1, "carrera_id": 1, "activo": True},
            ("FIS101",): {"_id": 32, "codigo_materia": "FIS101", "nombre": "Física I", "unidades_valorativas": 4,
                          "ciclo_plan": 1, "carrera_id": 1, "activo": True},
        }
        cambios, errores = self._analizar([_materia("MAT101"), _materia("FIS101", "Física I", uv=5), _materia("QUI101")],
                                          existentes)
        self.assertEqual(errores, [])
        self.assertEqual([c.tipo for c in cambios], ["igual", "cambio", "nueva"])
        self.assertEqual(cambios[1].registro_id, 32)
        self.assertEqual(cambios[1].diferencias, ["UV"])

    def test_los_errores_dicen_la_fila_y_el_campo(self):
        _, errores = self._analizar([_materia(carrera="XXX"), _materia("MAT2", uv=99), _materia("MAT3", nombre="")])
        self.assertEqual(len(errores), 3)
        self.assertTrue(errores[0].startswith("Fila 5:") and "«Carrera»" in errores[0] and "XXX" in errores[0])
        self.assertTrue(errores[1].startswith("Fila 6:") and "«UV»" in errores[1])
        self.assertTrue(errores[2].startswith("Fila 7:") and "«Nombre»" in errores[2])

    def test_un_codigo_repetido_en_el_archivo_se_rechaza(self):
        _, errores = self._analizar([_materia("MAT101"), _materia("mat101")])
        self.assertEqual(len(errores), 1)
        self.assertIn("ya aparece en la fila 5", errores[0])

    def test_activo_con_texto_raro_pide_si_o_no(self):
        _, errores = self._analizar([_materia(activo="quizás")])
        self.assertIn("Sí o No", errores[0])

    def test_otra_plantilla_no_se_acepta(self):
        with self.assertRaises(ErrorArchivo) as e:
            carga.analizar(MATERIAS, _con_cabecera(CATALOGOS["carreras"], ["SIS", "x", "D", 2020, 100, "Sí"]), REFS, {})
        self.assertIn("no coinciden con la plantilla", str(e.exception))

    def test_demasiadas_filas_piden_dividir_el_archivo(self):
        filas = [_materia(f"M{i}") for i in range(carga.FILAS_MAXIMAS + 1)]
        with self.assertRaises(ErrorArchivo) as e:
            self._analizar(filas)
        self.assertIn("Divídelo", str(e.exception))


class TestOtrosCatalogos(unittest.TestCase):
    def test_periodos_acepta_fechas_de_excel_y_solo_un_vigente(self):
        def fila(codigo, vigente):
            return [codigo, 2026, "I", datetime(2026, 2, 1), date(2026, 6, 30), vigente]
        cambios, errores = carga.analizar(PERIODOS, _con_cabecera(PERIODOS, fila("2026-I", "No"), fila("2026-II", "No")),
                                          REFS, {})
        self.assertEqual((errores, cambios[0].valores["fecha_inicio"]), ([], date(2026, 2, 1)))
        _, errores = carga.analizar(PERIODOS, _con_cabecera(PERIODOS, fila("2026-I", "Sí"), fila("2026-II", "Sí")), REFS, {})
        self.assertTrue(any("Solo un período puede quedar vigente" in e for e in errores))

    def test_periodos_con_fin_antes_del_inicio_se_rechaza(self):
        _, errores = carga.analizar(PERIODOS, _con_cabecera(
            PERIODOS, ["2026-I", 2026, "I", "2026-06-30", "2026-02-01", "No"]), REFS, {})
        self.assertIn("posterior", errores[0])

    def test_secciones_se_identifican_por_periodo_materia_y_numero(self):
        fila = ["2026-I", "MAT101", 1, "DOC-1", "MATUTINO", "a-1", 40]
        existentes = {(7, 11, 1): {"_id": 90, "periodo_id": 7, "materia_id": 11, "numero_seccion": 1, "docente_id": 5,
                                   "turno": "MATUTINO", "aula": "A-1", "cupo_maximo": 30}}
        cambios, errores = carga.analizar(SECCIONES, _con_cabecera(SECCIONES, fila), REFS, existentes)
        self.assertEqual(errores, [])
        self.assertEqual((cambios[0].tipo, cambios[0].registro_id, cambios[0].diferencias), ("cambio", 90, ["Cupo"]))

    def test_el_carnet_se_identifica_por_su_hash_y_no_viaja_en_la_confirmacion(self):
        fila = ["20240001", 2024, "SIS", "No", "REGULAR", "Sí"]
        cambios, errores = carga.analizar(ESTUDIANTES, _con_cabecera(ESTUDIANTES, fila), REFS, {})
        self.assertEqual(errores, [])
        self.assertEqual(cambios[0].carnet_hash, hash_carnet("20240001"))
        paquete = carga.empaquetar(ESTUDIANTES, cambios)
        self.assertNotIn("20240001", paquete)
        self.assertIn(hash_carnet("20240001"), paquete)

    def test_un_carnet_numerico_de_excel_es_el_mismo_que_en_texto(self):
        a, _ = carga.analizar(ESTUDIANTES, _con_cabecera(ESTUDIANTES, [20240001.0, 2024, "SIS", "No", "REGULAR", "Sí"]), REFS, {})
        b, _ = carga.analizar(ESTUDIANTES, _con_cabecera(ESTUDIANTES, ["20240001", 2024, "SIS", "No", "REGULAR", "Sí"]), REFS, {})
        self.assertEqual(a[0].carnet_hash, b[0].carnet_hash)


class TestConfirmacion(unittest.TestCase):
    def _cambios(self):
        cambios, _ = carga.analizar(MATERIAS, _con_cabecera(MATERIAS, _materia("MAT101"), _materia("FIS101")), REFS,
                                    {("FIS101",): {"_id": 32, "codigo_materia": "FIS101", "nombre": "x",
                                                   "unidades_valorativas": 4, "ciclo_plan": 1, "carrera_id": 1,
                                                   "activo": True}})
        return cambios

    def test_lo_revisado_vuelve_igual_al_confirmar(self):
        cambios = self._cambios()
        recuperados = carga.leer_carga(MATERIAS, carga.empaquetar(MATERIAS, cambios))
        self.assertEqual([(c.tipo, c.registro_id, c.valores) for c in recuperados],
                         [(c.tipo, c.registro_id, c.valores) for c in cambios])

    def test_las_filas_sin_cambio_no_viajan(self):
        igual = {("MAT101",): {"_id": 1, "codigo_materia": "MAT101", "nombre": "Matemática I",
                               "unidades_valorativas": 4, "ciclo_plan": 1, "carrera_id": 1, "activo": True}}
        cambios, _ = carga.analizar(MATERIAS, _con_cabecera(MATERIAS, _materia()), REFS, igual)
        self.assertEqual(json.loads(carga.empaquetar(MATERIAS, cambios)), [])

    def test_un_paquete_alterado_se_rechaza(self):
        paquete = json.loads(carga.empaquetar(MATERIAS, self._cambios()))
        casos = [
            "no es json",
            "[]",
            json.dumps([{"t": "borrar", "f": {}}]),
            json.dumps([{"t": "cambio", "i": "uno", "f": paquete[1]["f"]}]),           # sin id numérico
            json.dumps([dict(paquete[0], f=dict(paquete[0]["f"], unidades_valorativas="999"))]),   # fuera de rango
            json.dumps([dict(paquete[0], f=dict(paquete[0]["f"], codigo_materia=""))]),
        ]
        for caso in casos:
            with self.assertRaises(carga.ErrorCarga, msg=caso):
                carga.leer_carga(MATERIAS, caso)

    def test_un_hash_de_estudiante_invalido_se_rechaza(self):
        cambios, _ = carga.analizar(ESTUDIANTES, _con_cabecera(ESTUDIANTES, ["A1", 2024, "SIS", "No", "REGULAR", "Sí"]), REFS, {})
        paquete = json.loads(carga.empaquetar(ESTUDIANTES, cambios))
        paquete[0]["h"] = "no-es-un-hash"
        with self.assertRaises(carga.ErrorCarga):
            carga.leer_carga(ESTUDIANTES, json.dumps(paquete))


class TestAplicarCarga(unittest.TestCase):
    def test_crea_y_actualiza_en_un_solo_punto_de_guardado(self):
        db = mock.MagicMock()
        existente = SimpleNamespace(codigo_materia="FIS101", nombre="Viejo", unidades_valorativas=3, ciclo_plan=1,
                                    carrera_id=1, activo=True)
        db.get.return_value = existente
        cambios = carga.leer_carga(MATERIAS, json.dumps([
            {"t": "nueva", "n": 5, "i": None, "h": None,
             "f": {"codigo_materia": "QUI101", "nombre": "Química", "unidades_valorativas": "4", "ciclo_plan": "1",
                   "carrera_id": "1", "activo": "on"}},
            {"t": "cambio", "n": 6, "i": 32, "h": None,
             "f": {"codigo_materia": "FIS101", "nombre": "Física I", "unidades_valorativas": "5", "ciclo_plan": "1",
                   "carrera_id": "1", "activo": "on"}},
        ]))
        self.assertEqual(carga.aplicar_carga(db, MATERIAS, cambios), 2)
        db.begin_nested.assert_called_once()
        agregado = db.add.call_args.args[0]
        self.assertEqual((agregado.codigo_materia, agregado.nombre), ("QUI101", "Química"))
        self.assertEqual((existente.nombre, existente.unidades_valorativas), ("Física I", 5))

    def test_un_estudiante_nuevo_guarda_el_hash_y_nunca_el_carnet(self):
        db = mock.MagicMock()
        hash_ = hash_carnet("20240001")
        cambios = carga.leer_carga(ESTUDIANTES, json.dumps([
            {"t": "nueva", "n": 5, "i": None, "h": hash_,
             "f": {"anio_ingreso": "2024", "carrera_id": "1", "trabaja": "", "condicion_academica": "REGULAR",
                   "activo": "on"}}]))
        carga.aplicar_carga(db, ESTUDIANTES, cambios)
        self.assertEqual(db.add.call_args.args[0].carnet_hash, hash_)

    def test_un_registro_que_ya_no_existe_cancela_todo(self):
        db = mock.MagicMock()
        db.get.return_value = None
        cambios = carga.leer_carga(MATERIAS, json.dumps([
            {"t": "cambio", "n": 6, "i": 32, "h": None,
             "f": {"codigo_materia": "FIS101", "nombre": "Física I", "unidades_valorativas": "5", "ciclo_plan": "1",
                   "carrera_id": "1", "activo": "on"}}]))
        with self.assertRaises(carga.ErrorCarga):
            carga.aplicar_carga(db, MATERIAS, cambios)

    def test_el_periodo_vigente_desactiva_a_los_demas(self):
        db = mock.MagicMock()
        cambios = carga.leer_carga(PERIODOS, json.dumps([
            {"t": "nueva", "n": 5, "i": None, "h": None,
             "f": {"codigo_periodo": "2026-II", "anio": "2026", "ciclo_romano": "II", "fecha_inicio": "2026-08-01",
                   "fecha_fin": "2026-12-15", "activo": "on"}}]))
        carga.aplicar_carga(db, PERIODOS, cambios)
        sentencia = str(db.execute.call_args.args[0])
        self.assertIn("UPDATE", sentencia)
        self.assertIn("activo", sentencia)


class TestPlantilla(unittest.TestCase):
    def _libro(self, catalogo):
        from openpyxl import load_workbook
        campos = tuple(replace(c, fk=lambda db: [(1, "SIS · Ingeniería en Sistemas"), (2, "ADM · Administración")])
                       if c.tipo == "fk" else c for c in catalogo.campos)
        contenido = carga.construir_plantilla(mock.Mock(), replace(catalogo, campos=campos))
        return load_workbook(io.BytesIO(contenido))

    def test_la_primera_hoja_es_la_que_se_lee_y_trae_la_cabecera_del_catalogo(self):
        libro = self._libro(MATERIAS)
        self.assertEqual(libro.sheetnames, ["Datos", "Guía", "Referencia"])
        cabecera = [c.value for c in libro["Datos"][4]][:6]
        self.assertEqual(cabecera, carga.cabeceras(MATERIAS))

    def test_la_plantilla_descargada_se_vuelve_a_leer_sin_errores(self):
        libro = self._libro(MATERIAS)
        hoja = libro["Datos"]
        hoja.append([])
        hoja.cell(row=5, column=1, value="MAT101")
        for j, valor in enumerate(["Matemática I", 4, 1, "SIS", "Sí"], start=2):
            hoja.cell(row=5, column=j, value=valor)
        filas = [list(f) for f in hoja.iter_rows(values_only=True) if any(c is not None for c in f)]
        cambios, errores = carga.analizar(MATERIAS, filas, REFS, {})
        self.assertEqual((errores, len(cambios)), ([], 1))

    def test_la_referencia_lista_los_codigos_validos(self):
        libro = self._libro(MATERIAS)
        columna = [f[0] for f in libro["Referencia"].iter_rows(values_only=True)]
        self.assertEqual(columna, ["Carrera", "SIS · Ingeniería en Sistemas", "ADM · Administración"])

    def test_la_guia_explica_cada_campo(self):
        libro = self._libro(MATERIAS)
        guia = {f[0]: f for f in libro["Guía"].iter_rows(values_only=True)}
        self.assertIn("Número entero de 1 a 10", guia["UV"][2])
        self.assertIn("hoja Referencia", guia["Carrera"][2])
        self.assertEqual(guia["Activo"][1], "No")

    def test_todos_los_catalogos_tienen_clave_natural_y_plantilla(self):
        for slug, catalogo in CATALOGOS.items():
            self.assertIn(slug, carga.CLAVE_NATURAL)
            self.assertEqual(self._libro(catalogo).sheetnames[0], "Datos", slug)


if __name__ == "__main__":
    unittest.main()
